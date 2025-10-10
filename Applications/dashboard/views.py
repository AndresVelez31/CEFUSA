# Import necessary modules
from django.http import HttpResponse
from django.shortcuts import render, redirect
from django.contrib.auth.decorators import login_required
from django.contrib import messages
from collections import Counter
import io
import base64

import plotly.express as px
import plotly.graph_objects as go
import pandas as pd
from django import forms
from reportlab.pdfgen import canvas
from reportlab.lib.pagesizes import letter, A4
from reportlab.lib.utils import ImageReader
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Image, PageBreak, Table, TableStyle
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.units import inch
from reportlab.lib import colors
from datetime import datetime
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils.dataframe import dataframe_to_rows

from Applications.user.models import Player, Guardian
from Applications.payment.models import Payment 
from Applications.core.utils import user_can_access_dashboard
from datetime import datetime, timedelta

# This function is responsible for creating the logic behind of the dashboard
@login_required
def dashboard_view(request):
    """
    Vista del dashboard con datos y gráficos.
    Solo los usuarios Admin pueden acceder.
    """
    # Verificar si el usuario puede acceder al dashboard
    if not user_can_access_dashboard(request.user):
        messages.error(request, 'No tienes permisos para acceder al dashboard. Solo los administradores pueden verlo.')
        return redirect('homePage')
    players = Player.objects.all()
    guardians = Guardian.objects.all()
    payments = Payment.objects.all()

    # Player Charts
    # Pie chart: porcentaje de jugadores por tipo de documento
    tipo_doc_counts = {}
    total = players.count()
    for p in players:
        type = p.document_type if p.document_type else 'Sin documento'
        tipo_doc_counts[type] = tipo_doc_counts.get(type, 0) + 1

    fig_pie = px.pie(
        names=list(tipo_doc_counts.keys()),
        values=list(tipo_doc_counts.values()),
        title='jugadores por tipo de documento',
        hole=0.4,
        width=500,
        height=450
    )

    fig_pie.update_layout(
        title_text='Jugadores por tipo de documento',
        title_font=dict(size=24, family='Montserrat, sans-serif', color='black'),
        title_x=0.5  # Centra el título
    )
    
    chart_html_pie = fig_pie.to_html(full_html=False, include_plotlyjs='cdn')

    # Histogram of ages grouped by medical condition
    from datetime import date
    age_bins = [0, 5, 10, 15, 100]
    age_labels = ['0-5', '6-10', '11-15', 'Mayores a 15']
    today = date.today()
    ages = []
    conditions = []
    for p in players:
        if p.birth_date:
            age = today.year - p.birth_date.year - ((today.month, today.day) < (p.birth_date.month, p.birth_date.day))
            ages.append(age)
            conditions.append('Tiene condicion medica' if p.has_disease else 'No tiene condicion medica')
    df = pd.DataFrame({'age': ages, 'condition': conditions})
    df['grupos de edades'] = pd.cut(df['age'], bins=age_bins, labels=age_labels, right=True)
    fig_hist = px.histogram(
        df,
        x='grupos de edades',
        color='condition',
        barmode='group',
        title='Jugadores por Condición Médica',
        category_orders={'grupos de edades': age_labels},
        width=500,
        height=450
    )
    fig_hist.update_layout(
        title_text='Jugadores por Condición Médica',
        title_font=dict(size=24, family='Montserrat, sans-serif', color='black'),
        title_x=0.5,  # Centra el título
        yaxis_title='Cantidad'
    )
    chart_html_hist = fig_hist.to_html(full_html=False, include_plotlyjs=False)

    # Bar chart: students training in morning and afternoon
    sessions = [p.training_session if p.training_session else 'Sin jornada' for p in players]
    df_sessions = pd.DataFrame({'Jornada': sessions})
    session_counts = df_sessions['Jornada'].value_counts().reset_index()
    session_counts.columns = ['Jornada', 'Cantidad']
    fig_sessions = px.bar(
        session_counts,
        x='Jornada',
        y='Cantidad',
        title='Cantidad de estudiantes por jornada',
        color='Jornada',
        width=500,
        height=450
    )
    fig_sessions.update_layout(
        title_text='Cantidad de estudiantes por jornada',
        title_font=dict(size=20, family='Montserrat, sans-serif', color='black'),
        title_x=0.5,
        yaxis_title='Cantidad'
    )
    chart_html_sessions = fig_sessions.to_html(full_html=False, include_plotlyjs=False)

    # Chart: Students by category (birth year)
    categories = []
    for p in players:
        if p.birth_date:
            category = p.birth_date.year
            categories.append(str(category))
        else:
            categories.append('Sin fecha de nacimiento')
    
    df_categories = pd.DataFrame({'Categoría': categories})
    category_counts = df_categories['Categoría'].value_counts().reset_index()
    category_counts.columns = ['Categoría', 'Cantidad']
    
    # Ordenar por año de nacimiento (categoría) - filtrar solo años válidos
    valid_categories = category_counts[category_counts['Categoría'] != 'Sin fecha de nacimiento'].copy()
    invalid_categories = category_counts[category_counts['Categoría'] == 'Sin fecha de nacimiento'].copy()
    
    # Convertir a entero para ordenar correctamente y luego de vuelta a string
    if not valid_categories.empty:
        valid_categories['Categoría'] = valid_categories['Categoría'].astype(int)
        valid_categories = valid_categories.sort_values('Categoría')
        valid_categories['Categoría'] = valid_categories['Categoría'].astype(str)
    
    # Concatenar categorías válidas e inválidas
    category_counts = pd.concat([valid_categories, invalid_categories], ignore_index=True)
    
    fig_categories = px.bar(
        category_counts,
        x='Categoría',
        y='Cantidad',
        title='Cantidad de estudiantes por categoría',
        color='Categoría',
        width=500,
        height=450
    )
    fig_categories.update_layout(
        title_text='Cantidad de estudiantes por categoría',
        title_font=dict(size=20, family='Montserrat, sans-serif', color='black'),
        title_x=0.5,
        yaxis_title='Cantidad',
        xaxis_title='Categoría (Año de nacimiento)'
    )
    chart_html_eps = fig_categories.to_html(full_html=False, include_plotlyjs=False)

    # Guardian Charts
    # Guardian pie chart: IVA responsibility
    iva_types = [g.regime_type if g.regime_type else 'Sin dato' for g in guardians]
    iva_counts = pd.Series(iva_types).value_counts().reset_index()
    iva_counts.columns = ['Responsabilidad IVA', 'Cantidad']
    fig_guardian_iva = px.pie(
        iva_counts,
        names='Responsabilidad IVA',
        values='Cantidad',
        title='Acudientes responsables de IVA',
        hole=0.4,
        width=500,
        height=450
    )
    fig_guardian_iva.update_layout(
        title_text='Acudientes responsables de IVA',
        title_font=dict(size=24, family='Montserrat, sans-serif', color='black'),
        title_x=0.5
    )
    chart_html_guardian_iva = fig_guardian_iva.to_html(full_html=False, include_plotlyjs=False)

    # Guardian bar chart: most common cities
    guardian_cities = [g.city if g.city else 'Sin ciudad' for g in guardians]
    df_guardian_cities = pd.DataFrame({'Ciudad': guardian_cities})
    city_counts = df_guardian_cities['Ciudad'].value_counts().reset_index()
    city_counts.columns = ['Ciudad', 'Cantidad']
    fig_guardian_cities = px.bar(
        city_counts.head(10),
        x='Ciudad',
        y='Cantidad',
        title='Ciudades más frecuentes de acudientes',
        color='Ciudad',
        width=500,
        height=450
    )
    fig_guardian_cities.update_layout(
        title_text='Ciudades más frecuentes de acudientes',
        title_font=dict(size=20, family='Montserrat, sans-serif', color='black'),
        title_x=0.5,
        yaxis_title='Cantidad'
    )
    chart_html_guardian_cities = fig_guardian_cities.to_html(full_html=False, include_plotlyjs=False)

    # Payments Charts
    payments_df = pd.DataFrame(list(payments.values()))

    # 1. Ingresos mensuales del último año
    if not payments_df.empty:
        payments_df['date'] = pd.to_datetime(payments_df['date'])
        payments_df['amount'] = payments_df['amount'].astype(float)
        today = datetime.today()
        one_year_ago = today - timedelta(days=365)
        last_year_df = payments_df[(payments_df['date'] >= one_year_ago) & (payments_df['date'] <= today)]
        monthly_income = last_year_df.groupby(last_year_df['date'].dt.to_period('M')).agg({'amount': 'sum'}).reset_index()
        monthly_income['month'] = monthly_income['date'].astype(str)
        fig_payments_time = px.bar(
            monthly_income,
            x='month',
            y='amount',
            title='Ingresos mensuales a CEFUSA (último año)',
            color='month',
            width=500,
            height=450
        )
        fig_payments_time.update_layout(
            title_text='Ingresos mensuales a CEFUSA (último año)',
            title_font=dict(size=20, family='Montserrat, sans-serif', color='black'),
            title_x=0.5,
            yaxis_title='Valor'
        )
        chart_payments_time = fig_payments_time.to_html(full_html=False, include_plotlyjs='cdn')

        # 2. Cantidad de pagos a agrupados por sucursal
        cuentas_df = payments_df[payments_df['account'].isin(['5031', '5032'])]
        cuentas_group = cuentas_df.groupby(['account', 'branch']).size().reset_index(name='Cantidad')
        fig_cuentas_sucursal = px.bar(
            cuentas_group,
            x='branch',
            y='Cantidad',
            color='account',
            barmode='group',
            title='Pagos hechos por sucursal',
            width=500,
            height=450
        )
        fig_cuentas_sucursal.update_layout(
            title_text='Pagos hechos por sucursal',
            title_font=dict(size=20, family='Montserrat, sans-serif', color='black'),
            title_x=0.5,
            yaxis_title='Cantidad'
        )
        chart_cuentas_sucursal = fig_cuentas_sucursal.to_html(full_html=False, include_plotlyjs=False)

        # 3. Grafico de barras descripciones de pagos
        desc_group = payments_df['description'].value_counts().reset_index()
        desc_group.columns = ['Descripción', 'Cantidad']
        fig_desc = px.bar(
            desc_group.head(10),
            x='Descripción',
            y='Cantidad',
            color='Descripción',
            title='Cantidad de pagos por descripción',
            width=500,
            height=450
        )
        fig_desc.update_layout(
            title_text='Cantidad de pagos por descripción',
            title_font=dict(size=20, family='Montserrat, sans-serif', color='black'),
            title_x=0.5,
            yaxis_title='Cantidad'
        )
        chart_motivo = fig_desc.to_html(full_html=False, include_plotlyjs=False)

        # 4. Pagos por responsable y valor total en el mes
        if 'fk_responsible_id' in payments_df.columns:
            # Agrupar por responsable y mes, sumar valor
            payments_df['month'] = payments_df['date'].dt.to_period('M').astype(str)
            guardian_payments = payments_df.groupby(['fk_responsible_id', 'month']).agg({'amount': 'sum'}).reset_index()
            # Obtener nombres de responsables
            guardian_map = {g.id: f"{g.first_name} {g.last_name}" for g in guardians}
            guardian_payments['Responsable'] = guardian_payments['fk_responsible_id'].map(guardian_map)
            fig_ranking = px.bar(
                guardian_payments.sort_values('amount', ascending=False).head(10),
                x='Responsable',
                y='amount',
                color='month',
                title='Pagos realizados por responsables (valor total mensual)',
                width=500,
                height=450
            )
            fig_ranking.update_layout(
                title_text='Pagos realizados por responsables (valor total mensual)',
                title_font=dict(size=20, family='Montserrat, sans-serif', color='black'),
                title_x=0.5,
                yaxis_title='Valor total'
            )
            chart_ranking_acudientes = fig_ranking.to_html(full_html=False, include_plotlyjs=False)
        else:
            chart_ranking_acudientes = "<div>No hay datos de responsables disponibles.</div>"
    else:
        chart_payments_time = chart_cuentas_sucursal = chart_motivo = chart_ranking_acudientes = "<div>No hay datos de pagos disponibles.</div>"

    # We group all the charts in the context to be rendered in the template
    context = {
        'chart_pie': chart_html_pie,
        'chart_hist': chart_html_hist,
        'chart_sessions': chart_html_sessions,
        'chart_eps': chart_html_eps,
        'chart_guardian_iva': chart_html_guardian_iva,
        'chart_guardian_cities': chart_html_guardian_cities,
        'chart_payments_time': chart_payments_time,
        'chart_cuentas_sucursal': chart_cuentas_sucursal,
        'chart_motivo': chart_motivo,
        'chart_ranking_acudientes': chart_ranking_acudientes
    }

    return render(request, 'dashboard.html', context)


def generate_dashboard_pdf(request):
    """
    Genera un PDF con todos los datos estadísticos del dashboard en formato tabular
    """
    # Obtener los datos igual que en dashboard_view
    players = Player.objects.all()
    guardians = Guardian.objects.all()
    payments = Payment.objects.all()
    
    # Crear respuesta HTTP para PDF
    response = HttpResponse(content_type='application/pdf')
    response['Content-Disposition'] = f'attachment; filename="reporte_datos_cefusa_{datetime.now().strftime("%Y%m%d_%H%M%S")}.pdf"'
    
    # Crear el documento PDF
    doc = SimpleDocTemplate(response, pagesize=A4, leftMargin=50, rightMargin=50, topMargin=50, bottomMargin=50)
    story = []
    styles = getSampleStyleSheet()
    
    # Estilos personalizados
    title_style = ParagraphStyle(
        'CustomTitle',
        parent=styles['Heading1'],
        fontSize=18,
        spaceAfter=30,
        alignment=1,  # Centrado
        textColor='darkblue'
    )
    
    subtitle_style = ParagraphStyle(
        'CustomSubtitle',
        parent=styles['Heading2'],
        fontSize=14,
        spaceAfter=20,
        alignment=0,
        textColor='darkred'
    )
    
    # Título principal
    story.append(Paragraph("Reporte de Datos Estadísticos CEFUSA", title_style))
    story.append(Paragraph(f"Reporte generado el {datetime.now().strftime('%d/%m/%Y a las %H:%M')}", styles['Normal']))
    story.append(Spacer(1, 30))
    
    # Función helper para crear tablas
    def create_table(data, col_widths=None):
        table = Table(data, colWidths=col_widths)
        table.setStyle(TableStyle([
            ('BACKGROUND', (0, 0), (-1, 0), colors.grey),
            ('TEXTCOLOR', (0, 0), (-1, 0), colors.whitesmoke),
            ('ALIGN', (0, 0), (-1, -1), 'CENTER'),
            ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
            ('FONTSIZE', (0, 0), (-1, 0), 12),
            ('BOTTOMPADDING', (0, 0), (-1, 0), 12),
            ('BACKGROUND', (0, 1), (-1, -1), colors.beige),
            ('FONTNAME', (0, 1), (-1, -1), 'Helvetica'),
            ('FONTSIZE', (0, 1), (-1, -1), 10),
            ('GRID', (0, 0), (-1, -1), 1, colors.black)
        ]))
        return table
    
    try:
        # === RESUMEN GENERAL ===
        story.append(Paragraph("Resumen General", subtitle_style))
        
        summary_data = [
            ['Categoría', 'Cantidad Total'],
            ['Jugadores Registrados', str(players.count())],
            ['Acudientes Registrados', str(guardians.count())],
            ['Pagos Registrados', str(payments.count())]
        ]
        
        story.append(create_table(summary_data, [3*inch, 2*inch]))
        story.append(Spacer(1, 30))
        
        # === SECCIÓN JUGADORES ===
        story.append(Paragraph("Estadísticas Detalladas de Jugadores", subtitle_style))
        
        # 1. Tipos de documento
        story.append(Paragraph("Distribución por Tipo de Documento:", styles['Heading3']))
        tipo_doc_counts = {}
        for p in players:
            type_doc = p.document_type if p.document_type else 'Sin documento'
            tipo_doc_counts[type_doc] = tipo_doc_counts.get(type_doc, 0) + 1
        
        doc_data = [['Tipo de Documento', 'Cantidad', 'Porcentaje']]
        total_players = players.count()
        for doc_type, count in sorted(tipo_doc_counts.items()):
            percentage = f"{(count/total_players)*100:.1f}%" if total_players > 0 else "0%"
            doc_data.append([doc_type, str(count), percentage])
        
        story.append(create_table(doc_data, [2.5*inch, 1.5*inch, 1.5*inch]))
        story.append(Spacer(1, 20))
        
        # 2. Distribución por edades y condiciones médicas
        story.append(Paragraph("Distribución por Grupos de Edad y Condición Médica:", styles['Heading3']))
        from datetime import date
        age_bins = [0, 5, 10, 15, 100]
        age_labels = ['0-5 años', '6-10 años', '11-15 años', 'Mayores a 15 años']
        today = date.today()
        
        age_condition_data = {}
        for p in players:
            if p.birth_date:
                age = today.year - p.birth_date.year - ((today.month, today.day) < (p.birth_date.month, p.birth_date.day))
                # Determinar grupo de edad
                if age <= 5:
                    age_group = '0-5 años'
                elif age <= 10:
                    age_group = '6-10 años'
                elif age <= 15:
                    age_group = '11-15 años'
                else:
                    age_group = 'Mayores a 15 años'
                
                condition = 'Con condición médica' if p.has_disease else 'Sin condición médica'
                key = f"{age_group} - {condition}"
                age_condition_data[key] = age_condition_data.get(key, 0) + 1
        
        age_data = [['Grupo de Edad y Condición', 'Cantidad']]
        for key, count in sorted(age_condition_data.items()):
            age_data.append([key, str(count)])
        
        story.append(create_table(age_data, [4*inch, 1.5*inch]))
        story.append(Spacer(1, 20))
        
        # 3. Jornadas de entrenamiento
        story.append(Paragraph("Distribución por Jornada de Entrenamiento:", styles['Heading3']))
        session_counts = {}
        for p in players:
            session = p.training_session if p.training_session else 'Sin jornada asignada'
            session_counts[session] = session_counts.get(session, 0) + 1
        
        session_data = [['Jornada', 'Cantidad', 'Porcentaje']]
        for session, count in sorted(session_counts.items()):
            percentage = f"{(count/total_players)*100:.1f}%" if total_players > 0 else "0%"
            session_data.append([session, str(count), percentage])
        
        story.append(create_table(session_data, [2.5*inch, 1.5*inch, 1.5*inch]))
        story.append(Spacer(1, 20))
        
        # 4. Estudiantes por categoría (año de nacimiento)
        story.append(Paragraph("Distribución por Categoría (Año de Nacimiento):", styles['Heading3']))
        category_counts = {}
        for p in players:
            if p.birth_date:
                category = str(p.birth_date.year)
                category_counts[category] = category_counts.get(category, 0) + 1
            else:
                category_counts['Sin fecha de nacimiento'] = category_counts.get('Sin fecha de nacimiento', 0) + 1
        
        # Separar categorías válidas de inválidas y ordenar por año
        valid_categories = {k: v for k, v in category_counts.items() if k != 'Sin fecha de nacimiento'}
        invalid_categories = {k: v for k, v in category_counts.items() if k == 'Sin fecha de nacimiento'}
        
        # Ordenar categorías válidas por año
        sorted_valid = sorted(valid_categories.items(), key=lambda x: int(x[0]))
        sorted_invalid = list(invalid_categories.items())
        
        category_data = [['Categoría (Año)', 'Cantidad', 'Porcentaje']]
        for category, count in sorted_valid + sorted_invalid:
            percentage = f"{(count/total_players)*100:.1f}%" if total_players > 0 else "0%"
            category_data.append([category, str(count), percentage])
        
        story.append(create_table(category_data, [3*inch, 1*inch, 1*inch]))
        story.append(PageBreak())
        
        # === SECCIÓN ACUDIENTES ===
        story.append(Paragraph("Estadísticas Detalladas de Acudientes", subtitle_style))
        
        # 1. Responsabilidad de IVA
        story.append(Paragraph("Distribución por Responsabilidad de IVA:", styles['Heading3']))
        iva_counts = {}
        total_guardians = guardians.count()
        for g in guardians:
            iva_type = g.regime_type if g.regime_type else 'Sin dato'
            iva_counts[iva_type] = iva_counts.get(iva_type, 0) + 1
        
        iva_data = [['Tipo de Responsabilidad IVA', 'Cantidad', 'Porcentaje']]
        for iva_type, count in sorted(iva_counts.items()):
            percentage = f"{(count/total_guardians)*100:.1f}%" if total_guardians > 0 else "0%"
            iva_data.append([iva_type, str(count), percentage])
        
        story.append(create_table(iva_data, [3*inch, 1*inch, 1*inch]))
        story.append(Spacer(1, 20))
        
        # 2. Ciudades más frecuentes
        story.append(Paragraph("Distribución por Ciudad (Top 10):", styles['Heading3']))
        city_counts = {}
        for g in guardians:
            city = g.city if g.city else 'Sin ciudad'
            city_counts[city] = city_counts.get(city, 0) + 1
        
        # Ordenar por cantidad y tomar top 10
        sorted_cities = sorted(city_counts.items(), key=lambda x: x[1], reverse=True)[:10]
        
        city_data = [['Ciudad', 'Cantidad', 'Porcentaje']]
        for city, count in sorted_cities:
            percentage = f"{(count/total_guardians)*100:.1f}%" if total_guardians > 0 else "0%"
            city_data.append([city, str(count), percentage])
        
        story.append(create_table(city_data, [3*inch, 1*inch, 1*inch]))
        story.append(PageBreak())
        
        # === SECCIÓN PAGOS ===
        story.append(Paragraph("Estadísticas Detalladas de Pagos", subtitle_style))
        
        payments_df = pd.DataFrame(list(payments.values()))
        
        if not payments_df.empty and 'date' in payments_df.columns and 'amount' in payments_df.columns:
            payments_df['date'] = pd.to_datetime(payments_df['date'])
            payments_df['amount'] = payments_df['amount'].astype(float)
            
            # 1. Resumen financiero
            story.append(Paragraph("Resumen Financiero:", styles['Heading3']))
            total_amount = payments_df['amount'].sum()
            avg_amount = payments_df['amount'].mean()
            min_amount = payments_df['amount'].min()
            max_amount = payments_df['amount'].max()
            
            financial_data = [
                ['Concepto', 'Valor'],
                ['Total Recaudado', f"${total_amount:,.2f}"],
                ['Promedio por Pago', f"${avg_amount:,.2f}"],
                ['Pago Mínimo', f"${min_amount:,.2f}"],
                ['Pago Máximo', f"${max_amount:,.2f}"]
            ]
            
            story.append(create_table(financial_data, [3*inch, 2*inch]))
            story.append(Spacer(1, 20))
            
            # 2. Ingresos mensuales del último año
            story.append(Paragraph("Ingresos Mensuales (Último Año):", styles['Heading3']))
            today = datetime.today()
            one_year_ago = today - timedelta(days=365)
            last_year_df = payments_df[(payments_df['date'] >= one_year_ago) & (payments_df['date'] <= today)]
            
            if not last_year_df.empty:
                monthly_income = last_year_df.groupby(last_year_df['date'].dt.to_period('M')).agg({'amount': 'sum'}).reset_index()
                monthly_income['month'] = monthly_income['date'].astype(str)
                
                monthly_data = [['Mes', 'Ingresos']]
                for _, row in monthly_income.iterrows():
                    monthly_data.append([row['month'], f"${row['amount']:,.2f}"])
                
                story.append(create_table(monthly_data, [2*inch, 2*inch]))
                story.append(Spacer(1, 20))
            
            # 3. Pagos por sucursal (solo cuentas 5031 y 5032)
            story.append(Paragraph("Distribución de Pagos por Sucursal:", styles['Heading3']))
            if 'account' in payments_df.columns and 'branch' in payments_df.columns:
                cuentas_df = payments_df[payments_df['account'].isin(['5031', '5032'])]
                if not cuentas_df.empty:
                    branch_data = cuentas_df.groupby(['branch', 'account']).size().reset_index(name='cantidad')
                    
                    branch_table = [['Sucursal', 'Cuenta', 'Cantidad de Pagos']]
                    for _, row in branch_data.iterrows():
                        branch_table.append([str(row['branch']), str(row['account']), str(row['cantidad'])])
                    
                    story.append(create_table(branch_table, [2*inch, 1.5*inch, 1.5*inch]))
                    story.append(Spacer(1, 20))
            
            # 4. Descripciones de pagos más frecuentes
            story.append(Paragraph("Tipos de Pago Más Frecuentes (Top 10):", styles['Heading3']))
            if 'description' in payments_df.columns:
                desc_counts = payments_df['description'].value_counts().head(10)
                
                desc_data = [['Descripción', 'Cantidad', 'Porcentaje']]
                total_payments = len(payments_df)
                for desc, count in desc_counts.items():
                    percentage = f"{(count/total_payments)*100:.1f}%"
                    desc_data.append([str(desc), str(count), percentage])
                
                story.append(create_table(desc_data, [3*inch, 1*inch, 1*inch]))
        else:
            story.append(Paragraph("No hay datos de pagos disponibles para mostrar.", styles['Normal']))
        
        # Información del reporte
        story.append(PageBreak())
        story.append(Paragraph("Información del Reporte", subtitle_style))
        
        info_text = f"""
        <b>Fecha de generación:</b> {datetime.now().strftime('%d/%m/%Y %H:%M:%S')}<br/>
        <b>Total de registros procesados:</b><br/>
        • Jugadores: {players.count()}<br/>
        • Acudientes: {guardians.count()}<br/>
        • Pagos: {payments.count()}<br/>
        <b>Sistema:</b> CEFUSA - Centro de Formación de Fútbol<br/>
        <b>Tipo de reporte:</b> Estadísticas completas en formato tabular
        """
        
        story.append(Paragraph(info_text, styles['Normal']))
        
    except Exception as e:
        story.append(Paragraph(f"Error al procesar los datos: {str(e)}", styles['Normal']))
    
    # Construir PDF
    doc.build(story)
    return response


def generate_dashboard_excel(request):
    """
    Genera un archivo Excel con todas las estadísticas del dashboard en múltiples hojas
    """
    # Obtener los datos
    players = Player.objects.all()
    guardians = Guardian.objects.all()
    payments = Payment.objects.all()
    
    # Crear respuesta HTTP para Excel
    response = HttpResponse(content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
    response['Content-Disposition'] = f'attachment; filename="estadisticas_cefusa_{datetime.now().strftime("%Y%m%d_%H%M%S")}.xlsx"'
    
    # Crear workbook
    wb = openpyxl.Workbook()
    
    # Estilos
    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="366092", end_color="366092", fill_type="solid")
    center_alignment = Alignment(horizontal="center", vertical="center")
    
    def format_header(worksheet, start_row, end_row, start_col, end_col):
        """Función helper para formatear encabezados"""
        for row in range(start_row, end_row + 1):
            for col in range(start_col, end_col + 1):
                cell = worksheet.cell(row=row, column=col)
                cell.font = header_font
                cell.fill = header_fill
                cell.alignment = center_alignment
    
    # === HOJA 1: RESUMEN GENERAL ===
    ws_resumen = wb.active
    ws_resumen.title = "Resumen General"
    
    # Información del reporte
    ws_resumen.cell(row=1, column=1, value="REPORTE ESTADÍSTICAS CEFUSA")
    ws_resumen.cell(row=2, column=1, value=f"Generado el: {datetime.now().strftime('%d/%m/%Y %H:%M:%S')}")
    ws_resumen.cell(row=4, column=1, value="RESUMEN GENERAL")
    
    # Datos del resumen
    summary_data = [
        ["Categoría", "Cantidad Total"],
        ["Jugadores Registrados", players.count()],
        ["Acudientes Registrados", guardians.count()],
        ["Pagos Registrados", payments.count()]
    ]
    
    for i, row_data in enumerate(summary_data, start=6):
        for j, value in enumerate(row_data, start=1):
            ws_resumen.cell(row=i, column=j, value=value)
    
    format_header(ws_resumen, 6, 6, 1, 2)
    
    # === HOJA 2: JUGADORES ===
    ws_jugadores = wb.create_sheet("Jugadores")
    
    # Tipos de documento
    ws_jugadores.cell(row=1, column=1, value="ESTADÍSTICAS DE JUGADORES")
    ws_jugadores.cell(row=3, column=1, value="Distribución por Tipo de Documento")
    
    tipo_doc_counts = {}
    total_players = players.count()
    for p in players:
        type_doc = p.document_type if p.document_type else 'Sin documento'
        tipo_doc_counts[type_doc] = tipo_doc_counts.get(type_doc, 0) + 1
    
    doc_headers = ["Tipo de Documento", "Cantidad", "Porcentaje"]
    for j, header in enumerate(doc_headers, start=1):
        ws_jugadores.cell(row=5, column=j, value=header)
    
    row_num = 6
    for doc_type, count in sorted(tipo_doc_counts.items()):
        percentage = f"{(count/total_players)*100:.1f}%" if total_players > 0 else "0%"
        ws_jugadores.cell(row=row_num, column=1, value=doc_type)
        ws_jugadores.cell(row=row_num, column=2, value=count)
        ws_jugadores.cell(row=row_num, column=3, value=percentage)
        row_num += 1
    
    format_header(ws_jugadores, 5, 5, 1, 3)
    
    # Distribución por edades y condiciones médicas
    current_row = row_num + 2
    ws_jugadores.cell(row=current_row, column=1, value="Distribución por Grupos de Edad y Condición Médica")
    current_row += 2
    
    from datetime import date
    today = date.today()
    age_condition_data = {}
    
    for p in players:
        if p.birth_date:
            age = today.year - p.birth_date.year - ((today.month, today.day) < (p.birth_date.month, p.birth_date.day))
            if age <= 5:
                age_group = '0-5 años'
            elif age <= 10:
                age_group = '6-10 años'
            elif age <= 15:
                age_group = '11-15 años'
            else:
                age_group = 'Mayores a 15 años'
            
            condition = 'Con condición médica' if p.has_disease else 'Sin condición médica'
            key = f"{age_group} - {condition}"
            age_condition_data[key] = age_condition_data.get(key, 0) + 1
    
    ws_jugadores.cell(row=current_row, column=1, value="Grupo de Edad y Condición")
    ws_jugadores.cell(row=current_row, column=2, value="Cantidad")
    format_header(ws_jugadores, current_row, current_row, 1, 2)
    current_row += 1
    
    for key, count in sorted(age_condition_data.items()):
        ws_jugadores.cell(row=current_row, column=1, value=key)
        ws_jugadores.cell(row=current_row, column=2, value=count)
        current_row += 1
    
    # Jornadas de entrenamiento
    current_row += 2
    ws_jugadores.cell(row=current_row, column=1, value="Distribución por Jornada de Entrenamiento")
    current_row += 2
    
    session_counts = {}
    for p in players:
        session = p.training_session if p.training_session else 'Sin jornada asignada'
        session_counts[session] = session_counts.get(session, 0) + 1
    
    ws_jugadores.cell(row=current_row, column=1, value="Jornada")
    ws_jugadores.cell(row=current_row, column=2, value="Cantidad")
    ws_jugadores.cell(row=current_row, column=3, value="Porcentaje")
    format_header(ws_jugadores, current_row, current_row, 1, 3)
    current_row += 1
    
    for session, count in sorted(session_counts.items()):
        percentage = f"{(count/total_players)*100:.1f}%" if total_players > 0 else "0%"
        ws_jugadores.cell(row=current_row, column=1, value=session)
        ws_jugadores.cell(row=current_row, column=2, value=count)
        ws_jugadores.cell(row=current_row, column=3, value=percentage)
        current_row += 1
    
    # Estudiantes por categoría (año de nacimiento)
    current_row += 2
    ws_jugadores.cell(row=current_row, column=1, value="Estudiantes por Categoría (Año de Nacimiento)")
    current_row += 2
    
    category_counts = {}
    for p in players:
        if p.birth_date:
            category = str(p.birth_date.year)
            category_counts[category] = category_counts.get(category, 0) + 1
        else:
            category_counts['Sin fecha de nacimiento'] = category_counts.get('Sin fecha de nacimiento', 0) + 1
    
    # Separar categorías válidas de inválidas y ordenar por año
    valid_categories = {k: v for k, v in category_counts.items() if k != 'Sin fecha de nacimiento'}
    invalid_categories = {k: v for k, v in category_counts.items() if k == 'Sin fecha de nacimiento'}
    
    # Ordenar categorías válidas por año
    sorted_valid = sorted(valid_categories.items(), key=lambda x: int(x[0]))
    sorted_invalid = list(invalid_categories.items())
    sorted_categories = sorted_valid + sorted_invalid
    
    ws_jugadores.cell(row=current_row, column=1, value="Categoría (Año)")
    ws_jugadores.cell(row=current_row, column=2, value="Cantidad")
    ws_jugadores.cell(row=current_row, column=3, value="Porcentaje")
    format_header(ws_jugadores, current_row, current_row, 1, 3)
    current_row += 1
    
    for category, count in sorted_categories:
        percentage = f"{(count/total_players)*100:.1f}%" if total_players > 0 else "0%"
        ws_jugadores.cell(row=current_row, column=1, value=category)
        ws_jugadores.cell(row=current_row, column=2, value=count)
        ws_jugadores.cell(row=current_row, column=3, value=percentage)
        current_row += 1
    
    # === HOJA 3: ACUDIENTES ===
    ws_acudientes = wb.create_sheet("Acudientes")
    
    ws_acudientes.cell(row=1, column=1, value="ESTADÍSTICAS DE ACUDIENTES")
    ws_acudientes.cell(row=3, column=1, value="Distribución por Responsabilidad de IVA")
    
    total_guardians = guardians.count()
    iva_counts = {}
    for g in guardians:
        iva_type = g.regime_type if g.regime_type else 'Sin dato'
        iva_counts[iva_type] = iva_counts.get(iva_type, 0) + 1
    
    ws_acudientes.cell(row=5, column=1, value="Tipo de Responsabilidad IVA")
    ws_acudientes.cell(row=5, column=2, value="Cantidad")
    ws_acudientes.cell(row=5, column=3, value="Porcentaje")
    format_header(ws_acudientes, 5, 5, 1, 3)
    
    row_num = 6
    for iva_type, count in sorted(iva_counts.items()):
        percentage = f"{(count/total_guardians)*100:.1f}%" if total_guardians > 0 else "0%"
        ws_acudientes.cell(row=row_num, column=1, value=iva_type)
        ws_acudientes.cell(row=row_num, column=2, value=count)
        ws_acudientes.cell(row=row_num, column=3, value=percentage)
        row_num += 1
    
    # Ciudades más frecuentes
    current_row = row_num + 2
    ws_acudientes.cell(row=current_row, column=1, value="Distribución por Ciudad (Top 10)")
    current_row += 2
    
    city_counts = {}
    for g in guardians:
        city = g.city if g.city else 'Sin ciudad'
        city_counts[city] = city_counts.get(city, 0) + 1
    
    sorted_cities = sorted(city_counts.items(), key=lambda x: x[1], reverse=True)[:10]
    
    ws_acudientes.cell(row=current_row, column=1, value="Ciudad")
    ws_acudientes.cell(row=current_row, column=2, value="Cantidad")
    ws_acudientes.cell(row=current_row, column=3, value="Porcentaje")
    format_header(ws_acudientes, current_row, current_row, 1, 3)
    current_row += 1
    
    for city, count in sorted_cities:
        percentage = f"{(count/total_guardians)*100:.1f}%" if total_guardians > 0 else "0%"
        ws_acudientes.cell(row=current_row, column=1, value=city)
        ws_acudientes.cell(row=current_row, column=2, value=count)
        ws_acudientes.cell(row=current_row, column=3, value=percentage)
        current_row += 1
    
    # === HOJA 4: PAGOS ===
    ws_pagos = wb.create_sheet("Pagos")
    
    ws_pagos.cell(row=1, column=1, value="ESTADÍSTICAS DE PAGOS")
    
    payments_df = pd.DataFrame(list(payments.values()))
    
    if not payments_df.empty and 'date' in payments_df.columns and 'amount' in payments_df.columns:
        payments_df['date'] = pd.to_datetime(payments_df['date'])
        payments_df['amount'] = payments_df['amount'].astype(float)
        
        # Resumen financiero
        ws_pagos.cell(row=3, column=1, value="Resumen Financiero")
        
        total_amount = payments_df['amount'].sum()
        avg_amount = payments_df['amount'].mean()
        min_amount = payments_df['amount'].min()
        max_amount = payments_df['amount'].max()
        
        financial_data = [
            ["Concepto", "Valor"],
            ["Total Recaudado", f"${total_amount:,.2f}"],
            ["Promedio por Pago", f"${avg_amount:,.2f}"],
            ["Pago Mínimo", f"${min_amount:,.2f}"],
            ["Pago Máximo", f"${max_amount:,.2f}"]
        ]
        
        for i, row_data in enumerate(financial_data, start=5):
            for j, value in enumerate(row_data, start=1):
                ws_pagos.cell(row=i, column=j, value=value)
        
        format_header(ws_pagos, 5, 5, 1, 2)
        
        # Ingresos mensuales del último año
        current_row = 11
        ws_pagos.cell(row=current_row, column=1, value="Ingresos Mensuales (Último Año)")
        current_row += 2
        
        today = datetime.today()
        one_year_ago = today - timedelta(days=365)
        last_year_df = payments_df[(payments_df['date'] >= one_year_ago) & (payments_df['date'] <= today)]
        
        if not last_year_df.empty:
            monthly_income = last_year_df.groupby(last_year_df['date'].dt.to_period('M')).agg({'amount': 'sum'}).reset_index()
            monthly_income['month'] = monthly_income['date'].astype(str)
            
            ws_pagos.cell(row=current_row, column=1, value="Mes")
            ws_pagos.cell(row=current_row, column=2, value="Ingresos")
            format_header(ws_pagos, current_row, current_row, 1, 2)
            current_row += 1
            
            for _, row_data in monthly_income.iterrows():
                ws_pagos.cell(row=current_row, column=1, value=row_data['month'])
                ws_pagos.cell(row=current_row, column=2, value=f"${row_data['amount']:,.2f}")
                current_row += 1
        
        # Tipos de pago más frecuentes
        current_row += 2
        ws_pagos.cell(row=current_row, column=1, value="Tipos de Pago Más Frecuentes (Top 10)")
        current_row += 2
        
        if 'description' in payments_df.columns:
            desc_counts = payments_df['description'].value_counts().head(10)
            
            ws_pagos.cell(row=current_row, column=1, value="Descripción")
            ws_pagos.cell(row=current_row, column=2, value="Cantidad")
            ws_pagos.cell(row=current_row, column=3, value="Porcentaje")
            format_header(ws_pagos, current_row, current_row, 1, 3)
            current_row += 1
            
            total_payments = len(payments_df)
            for desc, count in desc_counts.items():
                percentage = f"{(count/total_payments)*100:.1f}%"
                ws_pagos.cell(row=current_row, column=1, value=str(desc))
                ws_pagos.cell(row=current_row, column=2, value=count)
                ws_pagos.cell(row=current_row, column=3, value=percentage)
                current_row += 1
    else:
        ws_pagos.cell(row=3, column=1, value="No hay datos de pagos disponibles")
    
    # Ajustar anchos de columnas en todas las hojas
    for ws in [ws_resumen, ws_jugadores, ws_acudientes, ws_pagos]:
        for column in ws.columns:
            max_length = 0
            column_letter = column[0].column_letter
            for cell in column:
                try:
                    if len(str(cell.value)) > max_length:
                        max_length = len(str(cell.value))
                except:
                    pass
            adjusted_width = min(max_length + 2, 50)
            ws.column_dimensions[column_letter].width = adjusted_width
    
    # Guardar el archivo en memoria
    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    
    response.write(output.getvalue())
    return response